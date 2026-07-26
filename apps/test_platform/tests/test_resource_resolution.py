from __future__ import annotations

import json
import tempfile
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from apps.test_platform.intent.builder import DefaultDesignBuilder
from apps.test_platform.intent.contracts import ReviewDecision, TestDesignRequest
from apps.test_platform.intent.prompt_builder import DefaultDesignPromptBuilder
from apps.test_platform.intent.service import TestDesignPipeline
from apps.test_platform.models import TestResourceProfile
from apps.test_platform.planning.compiler import TestPlanCompiler
from apps.test_platform.planning.contracts import PlanCandidate
from apps.test_platform.planning.resources import resolve_test_resources
from tests.test_fixtures_v4 import _approved_design
from .ui_asset_fixture import build_asset_database, procedure_payload


def _bundle():
    return _approved_design(
        "multichannel_initial_design_request.json",
        "multichannel_initial_design_candidate.json",
    )


class _ResourceGateway:
    def __init__(self, payload):
        self.payload = payload
        self.calls = 0

    def generate(self, messages, output_schema):
        self.calls += 1
        return output_schema.model_validate(self.payload)


def _single_ui_bundle(profile):
    candidate = {
        "title": "UI 首页验证",
        "objective": {"text": "验证首页显示就绪信息"},
        "in_scope": [{"text": "首页"}],
        "out_of_scope": [],
        "scenarios": [
            {
                "title": "首页就绪",
                "techniques": ["positive"],
                "requirement_ids": ["REQ-UI-1"],
                "operations": [{"text": "打开首页", "channel_hint": "ui"}],
                "expected_results": [
                    {
                        "text": "页面显示 Demo ready",
                        "after_operation_index": 1,
                        "channel_hint": "ui",
                    }
                ],
                "data_requirements": [],
                "state_impact": {
                    "impact": "read_only",
                    "rationale": {"text": "只读取页面"},
                },
            }
        ],
        "open_questions": [],
    }

    class Gateway:
        def generate(self, messages, output_schema):
            return output_schema.model_validate(candidate)

    request = TestDesignRequest.model_validate(
        {
            "schema_version": "test-design-request.v4",
            "request_id": "REQ-UI-RESOURCE-V4",
            "requirements": [
                {"requirement_id": "REQ-UI-1", "content": "首页显示 Demo ready"}
            ],
            "target": {
                "system_id": profile.system_id,
                "environment": profile.environment,
            },
            "selections": {
                "techniques": ["positive"],
                "allowed_channels": ["ui"],
                "required_channels": ["ui"],
                "knowledge_scope_ids": [],
            },
        }
    )
    pipeline = TestDesignPipeline(
        DefaultDesignBuilder(DefaultDesignPromptBuilder(), Gateway())
    )
    result = pipeline.generate(request)
    design, review = pipeline.review(
        result,
        decision=ReviewDecision.APPROVED,
        comments="UI 意图已核对",
    )
    return pipeline.build_approved_bundle(result, design, review)


class TestResourceResolutionTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.media.name)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        self.media.cleanup()

    def test_ui_reads_parameterized_procedure_from_selected_asset_database(self):
        source = Path(self.media.name) / "source.sqlite"
        first = procedure_payload(
            procedure_id="local.login",
            version=2,
            site="local.test",
        )
        library = build_asset_database(source, [first])
        profile = TestResourceProfile.objects.create(
            name="本地 UI",
            system_id="local-pages",
            environment="local",
            ui_procedure_database=SimpleUploadedFile(
                "local.sqlite",
                source.read_bytes(),
                content_type="application/vnd.sqlite3",
            ),
        )
        resolved = resolve_test_resources(profile, _bundle())

        ui = resolved.catalog.procedure_profiles[0]
        self.assertEqual(ui.site, "local.test")
        self.assertEqual(ui.library_id, library["library_id"])
        self.assertEqual(ui.library_hash, "sha256:" + library["library_hash"])
        login = ui.operations[0]
        self.assertEqual(login.procedure_id, "local.login")
        self.assertEqual(login.procedure_fingerprint, "sha256:" + library["rows"][0]["fingerprint"])
        self.assertEqual(login.procedure_parameters[0]["source_key"], "account_id")
        self.assertTrue(login.allowed_binding_refs)
        self.assertTrue(
            any(
                item.operation_ref == login.operation_ref
                for item in resolved.catalog.data_bindings
            )
        )
        self.assertEqual(
            resolved.runtime_config["procedure_asset_database"],
            profile.ui_procedure_database.path,
        )
        self.assertNotIn("browser_entry_points", resolved.runtime_config)

    def test_ui_resource_model_has_no_navigation_or_start_url_fields(self):
        field_names = {field.name for field in TestResourceProfile._meta.fields}
        self.assertIn("ui_procedure_database", field_names)
        self.assertTrue(
            {
                "ui_procedure_api_url",
                "ui_procedure_site",
                "ui_navigation_profile",
                "ui_start_url",
            }.isdisjoint(field_names)
        )

    def test_openapi_and_database_files_become_layer_two_catalog(self):
        openapi = {
            "openapi": "3.0.3",
            "paths": {
                "/health": {
                    "get": {
                        "operationId": "health.get",
                        "summary": "读取健康状态",
                        "responses": {"200": {"description": "ok"}},
                    }
                }
            },
        }
        queries = {
            "schema_version": "database-access-policy.v1",
            "database_schema": {
                "dialect": "sqlite",
                "tables": [
                    {
                        "name": "service_status",
                        "columns": [
                            {
                                "name": "status",
                                "data_type": "text",
                            }
                        ],
                    }
                ],
                "allowed_parameter_refs": [],
            },
        }
        profile = TestResourceProfile.objects.create(
            name="API 和数据库",
            api_openapi_file=SimpleUploadedFile(
                "openapi.json",
                json.dumps(openapi, ensure_ascii=False).encode("utf-8"),
            ),
            api_base_url="http://127.0.0.1:9000",
            database_query_file=SimpleUploadedFile(
                "queries.json",
                json.dumps(queries, ensure_ascii=False).encode("utf-8"),
            ),
            database_connection_ref="runtime.demo.db",
        )

        resolved = resolve_test_resources(profile, _bundle())

        self.assertEqual(
            resolved.catalog.available_executors,
            ["http_api", "database"],
        )
        self.assertEqual(
            resolved.catalog.http_operations[0].operation_ref,
            "api.operation.health.get",
        )
        self.assertEqual(
            resolved.runtime_config["base_urls"]["runtime.api.base"],
            "http://127.0.0.1:9000",
        )
        self.assertEqual(resolved.catalog.database_operations, [])
        self.assertEqual(
            resolved.catalog.database_schema.tables[0].name,
            "service_status",
        )
        self.assertNotIn("query_catalog", resolved.runtime_config)
        self.assertIn("database_schemas", resolved.runtime_config)

    def test_loose_resource_notes_are_normalized_once_and_cached(self):
        gateway = _ResourceGateway(
            {
                "api_operations": [
                    {
                        "name": "health",
                        "method": "GET",
                        "path": "/health/",
                        "description": "读取服务健康状态",
                        "parameters": [],
                    }
                ],
                "database_schema": {
                    "dialect": "sqlite",
                    "tables": [
                        {
                            "name": "service_status",
                            "description": "服务状态",
                            "columns": [
                                {
                                    "name": "status",
                                    "data_type": "text",
                                    "description": "当前状态",
                                }
                            ],
                        }
                    ],
                    "allowed_parameter_names": [],
                },
                "performance_profiles": [
                    {
                        "name": "health-smoke",
                        "description": "健康接口轻量负载",
                        "max_duration_seconds": 30,
                        "max_virtual_users": 5,
                        "target_api_operation": "health",
                        "metrics": [
                            {
                                "metric": "latency_ms",
                                "description": "响应延迟 p95",
                                "unit": "ms",
                                "percentile": "p95",
                            }
                        ],
                    }
                ],
            }
        )
        profile = TestResourceProfile.objects.create(
            name="宽松资源资料",
            system_id="commerce-hub",
            environment="matrix-local",
            api_asset_text="GET /health 返回服务状态",
            api_base_url="http://127.0.0.1:9000",
            database_asset_text="SQLite service_status 表包含 status 文本字段",
            database_connection_ref="runtime.demo.db",
            performance_asset_text="对 health 做 30 秒、5 用户负载，观察 p95 延迟",
        )

        first = resolve_test_resources(
            profile,
            _bundle(),
            resource_model_gateway=gateway,
        )

        self.assertEqual(gateway.calls, 1)
        self.assertEqual(
            first.catalog.available_executors,
            ["http_api", "database", "performance"],
        )
        self.assertEqual(first.catalog.http_operations[0].method, "GET")
        self.assertEqual(first.catalog.database_schema.tables[0].name, "service_status")
        self.assertEqual(first.catalog.performance_profiles[0].max_virtual_users, 5)
        self.assertEqual(
            first.runtime_config["performance_profiles"]["perf.health-smoke"]["inputs"],
            {"url": "http://127.0.0.1:9000/health/"},
        )

        profile.refresh_from_db()
        second = resolve_test_resources(profile, _bundle())
        self.assertEqual(second.catalog.content_hash, first.catalog.content_hash)
        self.assertEqual(gateway.calls, 1)

    def test_changed_loose_resource_requires_renormalization(self):
        profile = TestResourceProfile.objects.create(
            name="变化的宽松资料",
            api_asset_text="GET /health",
            api_base_url="http://127.0.0.1:9000",
            normalized_resource_data={
                "api_operations": [],
                "database_schema": None,
                "performance_profiles": [],
            },
            normalized_resource_source_hash="sha256:" + "0" * 64,
        )

        with self.assertRaisesRegex(ValueError, "重新生成执行计划"):
            resolve_test_resources(profile, _bundle())

    def test_ddl_file_is_accepted_as_loose_database_material(self):
        gateway = _ResourceGateway(
            {
                "api_operations": [],
                "database_schema": {
                    "dialect": "postgresql",
                    "tables": [
                        {
                            "name": "orders",
                            "description": "订单表",
                            "columns": [
                                {
                                    "name": "order_id",
                                    "data_type": "bigint",
                                    "description": "订单编号",
                                }
                            ],
                        }
                    ],
                    "allowed_parameter_names": [],
                },
                "performance_profiles": [],
            }
        )
        profile = TestResourceProfile.objects.create(
            name="DDL 数据库资料",
            database_query_file=SimpleUploadedFile(
                "schema.sql",
                b"CREATE TABLE orders (order_id BIGINT PRIMARY KEY);",
            ),
            database_connection_ref="runtime.orders.db",
        )

        resolved = resolve_test_resources(
            profile,
            _bundle(),
            resource_model_gateway=gateway,
        )

        self.assertEqual(gateway.calls, 1)
        self.assertEqual(resolved.catalog.database_schema.dialect, "postgresql")
        self.assertEqual(resolved.catalog.database_schema.tables[0].name, "orders")

    def test_database_access_policy_constrains_ai_sql_without_historical_queries(self):
        queries = {
            "schema_version": "database-access-policy.v1",
            "database_schema": {
                "dialect": "sqlite",
                "tables": [
                    {
                        "name": "service_status",
                        "description": "服务状态表",
                        "columns": [
                            {
                                "name": "status",
                                "data_type": "text",
                                "description": "服务状态",
                            }
                        ],
                    }
                ],
                "allowed_parameter_refs": ["runtime.service_id"],
            },
        }
        profile = TestResourceProfile.objects.create(
            name="AI SQL 数据库结构",
            database_query_file=SimpleUploadedFile(
                "database-v2.json",
                json.dumps(queries, ensure_ascii=False).encode("utf-8"),
            ),
            database_connection_ref="runtime.demo.db",
        )

        resolved = resolve_test_resources(profile, _bundle())

        self.assertEqual(resolved.catalog.available_executors, ["database"])
        self.assertEqual(resolved.catalog.database_operations, [])
        self.assertEqual(
            resolved.catalog.database_schema.tables[0].columns[0].name,
            "status",
        )
        self.assertEqual(
            resolved.runtime_config["database_schemas"]["runtime.demo.db"][
                "allowed_parameter_refs"
            ],
            ["runtime.service_id"],
        )

    def test_database_access_policy_rejects_historical_sql_content(self):
        policy_with_history = {
            "schema_version": "database-access-policy.v1",
            "database_schema": {
                "dialect": "sqlite",
                "tables": [
                    {
                        "name": "service_status",
                        "columns": [
                            {
                                "name": "status",
                                "data_type": "text",
                            }
                        ],
                    }
                ],
                "allowed_parameter_refs": [],
            },
            "queries": [
                {
                    "query_id": "db.service.status",
                    "sql": "SELECT status FROM service_status",
                }
            ],
        }
        profile = TestResourceProfile.objects.create(
            name="错误混入历史 SQL",
            database_query_file=SimpleUploadedFile(
                "invalid-policy.json",
                json.dumps(policy_with_history).encode("utf-8"),
            ),
            database_connection_ref="runtime.demo.db",
        )

        with self.assertRaisesRegex(ValueError, "历史 SQL.*业务知识库"):
            resolve_test_resources(profile, _bundle())

    def test_runtime_hash_changes_when_only_api_base_url_changes(self):
        openapi = {
            "openapi": "3.0.3",
            "paths": {
                "/health": {
                    "get": {
                        "operationId": "health.get",
                        "responses": {"200": {"description": "ok"}},
                    }
                }
            },
        }
        profile = TestResourceProfile.objects.create(
            name="API runtime drift",
            api_openapi_file=SimpleUploadedFile(
                "openapi.json", json.dumps(openapi).encode("utf-8")
            ),
            api_base_url="http://127.0.0.1:9001",
        )
        first = resolve_test_resources(profile, _bundle())
        profile.api_base_url = "http://127.0.0.1:9002"
        profile.save(update_fields=("api_base_url", "updated_at"))
        second = resolve_test_resources(profile, _bundle())

        self.assertEqual(first.catalog.content_hash, second.catalog.content_hash)
        self.assertNotEqual(first.runtime_config_hash, second.runtime_config_hash)

    def test_resource_protocols_reject_falsy_values_with_wrong_types(self):
        invalid_openapi = {
            "openapi": "3.0.3",
            "paths": {"/health": {"parameters": {}, "get": {"operationId": "health"}}},
        }
        api_profile = TestResourceProfile.objects.create(
            name="Invalid OpenAPI types",
            api_openapi_file=SimpleUploadedFile(
                "openapi.json", json.dumps(invalid_openapi).encode("utf-8")
            ),
            api_base_url="http://127.0.0.1:9000",
        )
        with self.assertRaisesRegex(ValueError, "parameters 必须是数组"):
            resolve_test_resources(api_profile, _bundle())

        invalid_queries = {
            "schema_version": "database-access-policy.v1",
            "database_schema": {
                "dialect": "sqlite",
                "tables": {},
            },
        }
        database_profile = TestResourceProfile.objects.create(
            name="Invalid database types",
            database_query_file=SimpleUploadedFile(
                "queries.json", json.dumps(invalid_queries).encode("utf-8")
            ),
            database_connection_ref="runtime.db",
        )
        with self.assertRaisesRegex(ValueError, "tables 必须是非空数组"):
            resolve_test_resources(database_profile, _bundle())

        invalid_performance = {
            "schema_version": "performance-profile-set.v1",
            "profiles": [
                {
                    "profile_ref": "perf.invalid",
                    "description": "invalid runtime type",
                    "driver_ref": "driver.http",
                    "state_effect": "read_only",
                    "max_duration_seconds": 60,
                    "max_virtual_users": 10,
                    "observables": [
                        {
                            "observable_ref": "observable.perf.invalid",
                            "description": "latency",
                            "metric": "latency_ms",
                        }
                    ],
                    "runtime": [],
                }
            ],
        }
        performance_profile = TestResourceProfile.objects.create(
            name="Invalid performance types",
            performance_profile_file=SimpleUploadedFile(
                "performance.json", json.dumps(invalid_performance).encode("utf-8")
            ),
        )
        with self.assertRaisesRegex(ValueError, "runtime 必须是对象"):
            resolve_test_resources(performance_profile, _bundle())

    def test_procedure_module_reaches_workbook_manifest(self):
        source = Path(self.media.name) / "plan-assets.sqlite"
        payload = procedure_payload(
            procedure_id="local.no-parameters",
            site="local.test",
        )
        payload["parameters"] = []
        payload["segments"][0]["items"] = [
            {"action": "click", "target": "Open"},
            {"action": "wait", "duration_ms": 0},
        ]
        library = build_asset_database(source, [payload])
        profile = TestResourceProfile.objects.create(
            name="本地 UI 计划",
            system_id="local-pages",
            environment="local",
            ui_procedure_database=SimpleUploadedFile(
                "plan-assets.sqlite",
                source.read_bytes(),
                content_type="application/vnd.sqlite3",
            ),
        )
        bundle = _single_ui_bundle(profile)
        resolved = resolve_test_resources(profile, bundle)

        scenario = bundle.design.scenarios[0]
        operation = scenario.operations[0]
        expected = scenario.expected_results[0]
        catalog_operation = resolved.catalog.procedure_profiles[0].operations[0]
        candidate = PlanCandidate.model_validate(
            {
                "flows": [
                    {
                        "scenario_id": scenario.scenario_id,
                        "stages": [
                            {
                                "executor_kind": "procedure_playwright",
                                "operations": [
                                    {
                                        "operation_id": operation.operation_id,
                                        "catalog_ref": catalog_operation.operation_ref,
                                    }
                                ],
                                "expected_results": [
                                    {
                                        "expected_result_id": expected.expected_result_id,
                                        "catalog_ref": catalog_operation.operation_ref,
                                        "observable_ref": next(
                                            item.observable_ref
                                            for item in resolved.catalog.procedure_profiles[0].observables
                                        ),
                                    }
                                ],
                            }
                        ],
                    }
                ],
                "open_questions": [],
            }
        )
        compiler = TestPlanCompiler()
        plan = compiler.build_draft(bundle, candidate, resolved.catalog)
        with tempfile.TemporaryDirectory() as output:
            compiled = compiler.compile(bundle, plan, resolved.catalog, output)
            self.assertTrue(compiled.validation.passed, compiled.validation.findings)
            workbook = next(Path(output).glob("**/case.xlsx"))
            manifest = next(Path(output).glob("**/manifest.json"))
            import openpyxl

            workbook_reader = openpyxl.load_workbook(workbook, read_only=True)
            rows = list(workbook_reader.active.values)
            workbook_reader.close()
            action_index = rows[0].index("Action")
            self.assertEqual(
                str(rows[1][action_index]),
                operation.text
                + " [procedure_id=local.no-parameters;version=1]",
            )
            manifest_value = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(manifest_value["capability_site"], "local.test")
            self.assertEqual(manifest_value["library_id"], library["library_id"])
            self.assertEqual(
                manifest_value["library_hash"],
                "sha256:" + library["library_hash"],
            )
            self.assertEqual(
                manifest_value["procedure_refs"],
                ["local.no-parameters@v1"],
            )
            self.assertEqual(
                manifest_value["procedure_calls"][0]["procedure_fingerprint"],
                "sha256:" + library["rows"][0]["fingerprint"],
            )
            self.assertNotIn("case_start_url", manifest_value)
            self.assertNotIn("navigation_profile", manifest_value)

    def test_performance_and_port_are_bounded_resources(self):
        profiles = {
            "schema_version": "performance-profile-set.v1",
            "profiles": [
                {
                    "profile_ref": "perf.health.smoke",
                    "description": "健康接口轻量性能检查",
                    "driver_ref": "driver.http",
                    "state_effect": "read_only",
                    "max_duration_seconds": 300,
                    "max_virtual_users": 100,
                    "observables": [
                        {
                            "observable_ref": "observable.perf.health-p95",
                            "description": "响应延迟 p95",
                            "metric": "latency_ms",
                            "unit": "ms",
                            "percentile": "p95",
                        }
                    ],
                    "runtime": {"target_url": "http://127.0.0.1:9000/health/"},
                }
            ],
        }
        profile = TestResourceProfile.objects.create(
            name="性能和端口",
            performance_profile_file=SimpleUploadedFile(
                "performance.json",
                json.dumps(profiles, ensure_ascii=False).encode("utf-8"),
            ),
            port_host="127.0.0.1",
            port_number=9000,
        )

        resolved = resolve_test_resources(profile, _bundle())

        self.assertEqual(
            resolved.catalog.available_executors,
            ["performance", "tcp_port"],
        )
        self.assertEqual(resolved.catalog.tcp_port_probes[0].port, 9000)
        self.assertEqual(
            resolved.runtime_config["network_hosts"]["runtime.port.host"],
            "127.0.0.1",
        )
        self.assertIn("perf.health.smoke", resolved.runtime_config["performance_profiles"])
        self.assertEqual(
            resolved.runtime_config["performance_profiles"]["perf.health.smoke"]["inputs"],
            {"url": "http://127.0.0.1:9000/health/"},
        )

    def test_http_performance_target_resolves_from_openapi_operation_id(self):
        openapi = {
            "openapi": "3.0.3",
            "paths": {
                "/health/": {
                    "get": {"operationId": "health", "summary": "健康检查"}
                }
            },
        }
        performance = {
            "schema_version": "performance-profile-set.v1",
            "profiles": [
                {
                    "profile_ref": "perf.health.smoke",
                    "description": "健康接口轻量性能检查",
                    "driver_ref": "driver.http",
                    "state_effect": "read_only",
                    "max_duration_seconds": 60,
                    "max_virtual_users": 10,
                    "observables": [
                        {
                            "observable_ref": "observable.perf.health-p95",
                            "description": "响应延迟 p95",
                            "metric": "latency_ms",
                            "unit": "ms",
                            "percentile": "p95",
                        }
                    ],
                    "runtime": {"target_operation_id": "health"},
                }
            ],
        }
        profile = TestResourceProfile.objects.create(
            name="API performance",
            api_openapi_file=SimpleUploadedFile(
                "openapi.json", json.dumps(openapi).encode("utf-8")
            ),
            api_base_url="http://127.0.0.1:8007",
            performance_profile_file=SimpleUploadedFile(
                "performance.json", json.dumps(performance).encode("utf-8")
            ),
        )

        resolved = resolve_test_resources(profile, _bundle())

        self.assertEqual(
            resolved.runtime_config["performance_profiles"]["perf.health.smoke"],
            {"inputs": {"url": "http://127.0.0.1:8007/health/"}},
        )
